"""``XlsxExtractor`` — two lockstep ``openpyxl`` read-only handles.

Scope & functionality
---------------------
Opens the workbook twice — ``data_only=False`` (formula text) and
``data_only=True`` (cached values) — and walks both worksheets'
``iter_rows()`` iterators together, row by row and cell by cell, so a
formula cell prefers its cached value and falls back to the formula text
itself when no cache exists (never an empty cell — the fix for the
audited ``formula_without_cached_value`` gap). Row and column caps are
enforced by breaking the iteration, never by materialising a sheet first
(``list(ws.iter_rows())``), so a 200k-row sheet stays streaming.

Merge spans come from ``xlsx.merged.MergedRangeReader`` (``read_only``
leaves ``ws.merged_cells`` empty); only the origin cell of a merge is
emitted, carrying its ``rowspan``/``colspan`` — every other cell the merge
covers (to its right in the same row, or in the rows below) is *omitted*
from ``Row.cells`` entirely, the way ``GfmRenderer``'s grid expansion
already expects (it reconstructs the covered slots from the span, so a
placeholder cell there would double-count a column and misalign the row).
The header row is decided by
``xlsx.header.HeaderRowHeuristic``; cell text by ``xlsx.cells.CellFormatter``.
Each sheet is isolated in its own try/except: one bad sheet becomes a
``UnitFailure`` and extraction continues; a workbook whose every sheet
fails (or that doesn't open as a workbook at all) becomes ``CorruptFile``.
"""

from __future__ import annotations

import dataclasses
import zipfile
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]
from openpyxl.workbook.workbook import Workbook  # type: ignore[import-untyped]

from evoq_ingestion.core.codes import WarningCode
from evoq_ingestion.core.limits import Deadline, Limits
from evoq_ingestion.extractors.base import BaseExtractor, CorruptFile, EngineInfo, ExtractionError
from evoq_ingestion.extractors.gate import FileFacts
from evoq_ingestion.extractors.xlsx.cells import CellFormatter
from evoq_ingestion.extractors.xlsx.header import HeaderRowHeuristic
from evoq_ingestion.extractors.xlsx.merged import MergedRangeReader, MergeMap
from evoq_ingestion.model.blocks import Cell, Row, SheetBreak, Table
from evoq_ingestion.model.document import Document, SheetInfo

_ROW_CHECK_INTERVAL = 2000


class XlsxExtractor(BaseExtractor):
    """``openpyxl`` engine: one ``--- Sheet: {name} ---`` marker + table per sheet."""

    info = EngineInfo("openpyxl", "2")
    supports = frozenset({".xlsx"})

    def classify_error(self, exc: Exception, path: Path) -> ExtractionError:
        """Any unclassified zip/openpyxl failure means the container is unreadable."""
        if isinstance(exc, zipfile.BadZipFile | InvalidFileException | KeyError):
            return CorruptFile(f"corrupt xlsx: {path.name}", detail=str(exc))
        return super().classify_error(exc, path)

    def _parse(self, path: Path, facts: FileFacts, deadline: Deadline) -> Document:
        """Build the ``Document`` for one workbook, degrading sheet by sheet."""
        doc = Document(title=path.stem)
        limits = self.options.limits
        wb_formulas = self._open(path, data_only=False)
        try:
            wb_values = self._open(path, data_only=True)
        except ExtractionError:
            wb_formulas.close()
            raise
        try:
            sheet_names = list(wb_formulas.sheetnames)
            merge_reader = MergedRangeReader(path, limits)
            any_ok = False
            for name in sheet_names:
                deadline.check()
                ok = self._process_sheet(
                    doc, wb_formulas[name], wb_values[name], name, merge_reader, limits, deadline
                )
                any_ok = any_ok or ok
            if sheet_names and not any_ok:
                raise CorruptFile(f"every sheet failed in {path.name}")
        finally:
            wb_formulas.close()
            wb_values.close()
        return doc

    @staticmethod
    def _open(path: Path, *, data_only: bool) -> Workbook:
        try:
            return openpyxl.load_workbook(str(path), read_only=True, data_only=data_only)
        except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
            raise CorruptFile(f"corrupt xlsx: {path.name}", detail=str(exc)) from exc

    def _process_sheet(
        self,
        doc: Document,
        ws_formulas: Any,  # noqa: ANN401 - untyped openpyxl read-only worksheet
        ws_values: Any,  # noqa: ANN401
        name: str,
        merge_reader: MergedRangeReader,
        limits: Limits,
        deadline: Deadline,
    ) -> bool:
        """Extract one sheet; ``True`` unless this sheet failed outright."""
        hidden = ws_formulas.sheet_state != "visible"
        if hidden and self.options.xlsx_hidden_sheets == "skip":
            doc.add(SheetBreak(name=name, hidden=True))
            doc.add_warning(WarningCode.hidden_sheet_skipped, name)
            doc.units.sheets.append(SheetInfo(name=name, rows=0, cols=0, hidden=True))
            return True
        try:
            table, n_rows, n_cols = self._read_sheet(
                doc, ws_formulas, ws_values, name, merge_reader, limits, deadline
            )
        except Exception as exc:  # noqa: BLE001 - isolate one bad sheet from the rest
            doc.fail_unit(
                "sheet", name, WarningCode.extraction_failed, f"{type(exc).__name__}: {exc}"
            )
            doc.units.sheets.append(SheetInfo(name=name, rows=0, cols=0, hidden=hidden))
            return False
        doc.add(SheetBreak(name=name, hidden=hidden))
        if hidden:
            doc.add_warning(WarningCode.hidden_sheet, name)
        if table is not None:
            doc.add(table)
        doc.units.sheets.append(SheetInfo(name=name, rows=n_rows, cols=n_cols, hidden=hidden))
        return True

    def _read_sheet(
        self,
        doc: Document,
        ws_formulas: Any,  # noqa: ANN401
        ws_values: Any,  # noqa: ANN401
        name: str,
        merge_reader: MergedRangeReader,
        limits: Limits,
        deadline: Deadline,
    ) -> tuple[Table | None, int, int]:
        merges = merge_reader.read(name)
        if merges.skipped:
            doc.add_warning(WarningCode.merged_scan_skipped, name)

        rows, header_row, next_row, stats = self._stream_rows(
            ws_formulas, ws_values, limits, merges, deadline
        )
        rows_truncated, cols_truncated, formula_fallbacks, max_col_seen = stats

        if rows_truncated:
            doc.add_warning(WarningCode.rows_truncated, f"{name}:{limits.max_rows_per_sheet}")
        if cols_truncated:
            doc.add_warning(WarningCode.cols_truncated, f"{name}:{limits.max_cols_per_sheet}")
        if formula_fallbacks:
            doc.add_warning(WarningCode.formula_without_cached_value, f"{name}:{formula_fallbacks}")

        if not rows:
            return None, 0, 0

        has_header = header_row is not None and HeaderRowHeuristic.looks_like_header(
            header_row, next_row
        )
        if header_row is not None and not has_header:
            doc.add_warning(WarningCode.no_header_row, name)
        header_rows = 1 if has_header else 0
        if has_header:
            rows[0] = Row(cells=tuple(dataclasses.replace(c, header=True) for c in rows[0].cells))

        data_rows = max(len(rows) - header_rows, 0)
        return Table(rows=tuple(rows), header_rows=header_rows, name=name), data_rows, max_col_seen

    @staticmethod
    def _stream_rows(
        ws_formulas: Any,  # noqa: ANN401
        ws_values: Any,  # noqa: ANN401
        limits: Limits,
        merges: MergeMap,
        deadline: Deadline,
    ) -> tuple[list[Row], list[Any] | None, list[Any] | None, tuple[bool, bool, int, int]]:
        """Stream both handles in lockstep, breaking at the row/column caps."""
        max_rows = limits.max_rows_per_sheet
        max_cols = limits.max_cols_per_sheet
        rows: list[Row] = []
        header_row: list[Any] | None = None
        next_row: list[Any] | None = None
        rows_truncated = False
        cols_truncated = False
        formula_fallbacks = 0
        max_col_seen = 0

        # `strict=True`: the two handles read the same file, so a length
        # mismatch is a real defect worth surfacing (caught by the caller's
        # per-sheet try/except) rather than silently misaligning rows.
        paired = zip(ws_formulas.iter_rows(), ws_values.iter_rows(), strict=True)
        for row_idx, (row_f, row_v) in enumerate(paired, start=1):
            if row_idx > max_rows:
                rows_truncated = True
                break
            if row_idx % _ROW_CHECK_INTERVAL == 0:
                deadline.check()
            cells, probe_values, col_count, row_cols_truncated, row_fallbacks = _read_row(
                row_f, row_v, row_idx, max_cols, merges
            )
            cols_truncated = cols_truncated or row_cols_truncated
            formula_fallbacks += row_fallbacks
            max_col_seen = max(max_col_seen, col_count)
            rows.append(Row(cells=tuple(cells)))
            if header_row is None:
                if any(v not in (None, "") for v in probe_values):
                    header_row = probe_values
            elif next_row is None:
                next_row = probe_values

        stats = (rows_truncated, cols_truncated, formula_fallbacks, max_col_seen)
        return rows, header_row, next_row, stats


def _read_row(
    row_f: tuple[Any, ...],
    row_v: tuple[Any, ...],
    row_idx: int,
    max_cols: int,
    merges: MergeMap,
) -> tuple[list[Cell], list[Any], int, bool, int]:
    """One row's cells, its header-probe values, the column count, truncation, fallbacks.

    ``probe_values`` is position-aligned to the sheet's physical columns
    (a merge-covered slot is ``None`` there) so ``HeaderRowHeuristic`` sees
    the row's true density; ``cells`` omits covered slots entirely, since
    the origin cell's ``colspan``/``rowspan`` already accounts for them.
    """
    cells: list[Cell] = []
    probe_values: list[Any] = []
    truncated = False
    fallbacks = 0
    col_idx = 0
    for col_idx, (cell_f, cell_v) in enumerate(zip(row_f, row_v, strict=True), start=1):
        if col_idx > max_cols:
            truncated = True
            col_idx -= 1  # report the capped width, not the first overflow column
            break
        value, used_fallback = _resolve_value(cell_f, cell_v)
        fallbacks += 1 if used_fallback else 0
        covered = merges.is_covered(row_idx, col_idx)
        probe_values.append(None if covered else value)
        if covered:
            continue
        span = merges.span_at(row_idx, col_idx)
        if span is not None:
            cells.append(Cell(text=CellFormatter.format(value), rowspan=span[0], colspan=span[1]))
        else:
            cells.append(Cell(text=CellFormatter.format(value)))
    return cells, probe_values, col_idx, truncated, fallbacks


def _resolve_value(cell_f: Any, cell_v: Any) -> tuple[Any, bool]:  # noqa: ANN401
    """Cached value wins; a formula with no cache falls back to its formula text."""
    if getattr(cell_f, "data_type", None) == "f":
        if cell_v.value is not None:
            return cell_v.value, False
        return cell_f.value, True
    return cell_v.value, False
