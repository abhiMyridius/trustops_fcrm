#!/usr/bin/env python
"""Deterministic generator for the shared AEO test corpus (plan §3.3).

Scope & functionality
---------------------
Writes every fixture referenced by ``scenarios.yaml`` into ``generated/``
(and the chunk-only Markdown inputs into ``ctx/``) from code alone — no
network, no clock, no machine-specific state.  Running the script twice
must produce byte-identical files; ``--check`` proves it by regenerating
into a temporary directory and diffing against what is on disk.

Determinism recipe (keep it when adding fixtures):

* one seed, ``SEED = 20260911``; every fixture draws from its **own**
  ``random.Random`` derived from ``sha256(f"{SEED}:{fixture_id}")`` so adding
  a fixture never shifts another one's bytes;
* PDF: ``reportlab`` ``canvas.Canvas(..., invariant=1)`` (fixed CreationDate
  and document ID) with fixed metadata; only the 14 standard fonts or the
  built-in CID font ``STSong-Light`` (nothing embedded from the host);
* OOXML (``python-docx`` / ``openpyxl`` / ``python-pptx``): fixed core
  properties, then ``_normalise_zip`` rewrites the package with every
  ``ZipInfo.date_time = (1980, 1, 1, 0, 0, 0)``, a fixed compression level
  and a deterministic entry order;
* HTML / ADF / Markdown / text: literal content in this file;
* the 201 MiB sparse file (EDGE-01) is created with ``seek`` + one byte,
  lives under ``generated/edge/`` (git-ignored) and is only produced with
  ``--edge``.

Usage::

    python generate_corpus.py                 # (re)generate generated/ + ctx/
    python generate_corpus.py --check         # regenerate to a temp dir, diff bytes, exit 1 on drift
    python generate_corpus.py --edge          # also create the on-demand EDGE-01 sparse file
    python generate_corpus.py --only 'PDF-*'  # subset
    python generate_corpus.py --list          # print the fixture registry

Requirements: ``requirements-corpus.txt`` (reportlab, python-docx, openpyxl,
python-pptx, Pillow, PyYAML).
"""

from __future__ import annotations

import argparse
import base64
import fnmatch
import hashlib
import io
import json
import os
import random
import re
import sys
import tempfile
import zipfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path

# --------------------------------------------------------------------------- constants

SEED = 20260911
CORPUS_ROOT = Path(__file__).resolve().parent
GENERATED_DIRNAME = "generated"
EDGE_DIRNAME = "edge"
CTX_DIRNAME = "ctx"

FIXED_DT = datetime(2026, 9, 11, 0, 0, 0, tzinfo=UTC)
FIXED_DT_NAIVE = datetime(2026, 9, 11, 0, 0, 0)
ZIP_DATE = (1980, 1, 1, 0, 0, 0)
ZIP_LEVEL = 9
AUTHOR = "EVOQ corpus generator"
MAX_COMMITTED_BYTES = 300 * 1024

SPARSE_BYTES = 201 * 1024 * 1024  # EDGE-01: one MiB over the 200 MiB cap


# --------------------------------------------------------------------------- helpers


def _rng(tag: str) -> random.Random:
    """A per-fixture RNG: stable across runs, independent of other fixtures."""
    digest = hashlib.sha256(f"{SEED}:{tag}".encode()).hexdigest()
    return random.Random(int(digest[:16], 16))


_GENERAL_WORDS = [
    "system",
    "service",
    "request",
    "response",
    "ledger",
    "account",
    "posting",
    "balance",
    "audit",
    "policy",
    "operator",
    "tenant",
    "session",
    "token",
    "schedule",
    "batch",
    "report",
    "threshold",
    "retry",
    "timeout",
    "record",
    "version",
    "snapshot",
    "manifest",
    "index",
    "vector",
    "document",
    "chunk",
    "heading",
    "table",
    "requirement",
    "approval",
    "reviewer",
    "owner",
    "deadline",
    "metric",
    "alert",
    "queue",
    "worker",
    "handler",
]

_DATABASE_WORDS = [
    "database",
    "table",
    "index",
    "query",
    "transaction",
    "commit",
    "rollback",
    "replica",
    "shard",
    "partition",
    "schema",
    "column",
    "row",
    "primary",
    "key",
    "foreign",
    "constraint",
    "latency",
    "throughput",
    "cache",
    "buffer",
    "checkpoint",
    "journal",
    "vacuum",
    "statistics",
    "planner",
    "join",
    "scan",
    "lock",
    "deadlock",
    "isolation",
]

_GARDEN_WORDS = [
    "garden",
    "soil",
    "compost",
    "seedling",
    "mulch",
    "trellis",
    "pruning",
    "blossom",
    "harvest",
    "orchard",
    "irrigation",
    "shade",
    "perennial",
    "pollinator",
    "hedge",
    "lawn",
    "tomato",
    "basil",
    "rosemary",
    "greenhouse",
    "frost",
    "drainage",
    "sunlight",
    "fertiliser",
    "weeding",
    "raised",
    "bed",
    "sprout",
    "canopy",
    "bloom",
]

_FILLERS = ["the", "a", "each", "every", "this", "that", "its", "their", "any", "no"]
_VERBS = [
    "shall",
    "must",
    "may",
    "records",
    "returns",
    "validates",
    "rejects",
    "stores",
    "emits",
    "checks",
]


def _sentence(rng: random.Random, words: list[str], n: int | None = None) -> str:
    n = n or rng.randint(7, 15)
    out: list[str] = []
    for i in range(n):
        if i % 4 == 1:
            out.append(rng.choice(_VERBS))
        elif i % 5 == 3:
            out.append(rng.choice(_FILLERS))
        else:
            out.append(rng.choice(words))
    text = " ".join(out)
    return text[0].upper() + text[1:] + "."


def _paragraph(
    rng: random.Random, words: list[str] | None = None, sentences: int | None = None
) -> str:
    words = words or _GENERAL_WORDS
    return " ".join(
        _sentence(rng, words) for _ in range(sentences or rng.randint(3, 6))
    )


def _wrap(text: str, width: int) -> list[str]:
    import textwrap

    return textwrap.wrap(text, width=width) or [""]


def _normalise_zip(
    data: bytes,
    patches: dict[str, Callable[[bytes], bytes]] | None = None,
    extra: dict[str, bytes] | None = None,
) -> bytes:
    """Rewrite an OOXML/zip package deterministically.

    * every entry gets ``date_time == (1980, 1, 1, 0, 0, 0)``;
    * ``[Content_Types].xml`` first (OOXML convention), then sorted names;
    * one compression method and level; fixed ``create_system``/``external_attr``.
    ``patches`` transform existing parts, ``extra`` adds new ones.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as src:
        entries = {info.filename: src.read(info) for info in src.infolist()}
    for name, fn in (patches or {}).items():
        if name not in entries:
            raise KeyError(f"zip patch target missing: {name}")
        entries[name] = fn(entries[name])
    entries.update(extra or {})
    names = sorted(entries)
    if "[Content_Types].xml" in names:
        names.remove("[Content_Types].xml")
        names.insert(0, "[Content_Types].xml")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for name in names:
            zi = zipfile.ZipInfo(name, date_time=ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 3
            zi.external_attr = 0o644 << 16
            dst.writestr(zi, entries[name], compresslevel=ZIP_LEVEL)
    return out.getvalue()


def _sub_once(pattern: str, repl: str, text: str, flags: int = 0) -> str:
    new, n = re.subn(pattern, repl, text, count=1, flags=flags)
    if n != 1:
        raise ValueError(f"pattern not found for patch: {pattern!r}")
    return new


# --------------------------------------------------------------------------- PDF helpers


class _PdfWriter:
    """Minimal flowing-text writer over a reportlab canvas (deterministic, invariant)."""

    def __init__(self, path: Path, title: str, *, encrypt=None) -> None:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas as rl_canvas

        kwargs = {"encrypt": encrypt} if encrypt is not None else {}
        self.c = rl_canvas.Canvas(
            str(path), pagesize=A4, invariant=1, pageCompression=1, **kwargs
        )
        self.c.setAuthor(AUTHOR)
        self.c.setTitle(title)
        self.c.setSubject("AEO deterministic test corpus")
        self.c.setCreator("generate_corpus.py")
        self.c.setKeywords("evoq corpus fixture")
        self.width, self.height = A4
        self.margin = 56.0
        self.y = self.height - self.margin
        self.page = 1
        self.header: str | None = None
        self.footer: str | None = None  # may contain {page} / {pages}
        self.pages_total: int | None = None
        self._page_started = False

    # -- page management
    def _ensure_page(self) -> None:
        if not self._page_started:
            self._draw_chrome()
            self._page_started = True

    def _draw_chrome(self) -> None:
        if self.header:
            self.c.setFont("Helvetica", 8)
            self.c.drawString(self.margin, self.height - 30, self.header)
        if self.footer:
            self.c.setFont("Helvetica", 8)
            txt = self.footer.format(page=self.page, pages=self.pages_total or "?")
            self.c.drawString(self.margin, 24, txt)

    def new_page(self) -> None:
        self._ensure_page()
        self.c.showPage()
        self.page += 1
        self.y = self.height - self.margin
        self._page_started = False

    def _need(self, h: float) -> None:
        self._ensure_page()
        if self.y - h < self.margin:
            self.new_page()
            self._ensure_page()

    # -- content
    def heading(
        self, text: str, size: float, *, bold: bool = True, space_before: float = 10
    ) -> None:
        self.y -= space_before
        self._need(size * 1.6)
        self.c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        self.c.drawString(self.margin, self.y - size, text)
        self.y -= size * 1.6

    def para(
        self,
        text: str,
        size: float = 10,
        *,
        font: str = "Helvetica",
        x: float | None = None,
        width: float | None = None,
        leading: float | None = None,
        chars: int | None = None,
    ) -> None:
        x = self.margin if x is None else x
        width = (self.width - 2 * self.margin) if width is None else width
        leading = leading or size * 1.3
        chars = chars or max(20, int(width / (size * 0.5)))
        for line in _wrap(text, chars):
            self._need(leading)
            self.c.setFont(font, size)
            self.c.drawString(x, self.y - size, line)
            self.y -= leading
        self.y -= leading * 0.6

    def table(
        self,
        rows: list[list[str]],
        col_widths: list[float],
        *,
        size: float = 9,
        row_h: float = 16,
        caption: str | None = None,
    ) -> None:
        if caption:
            self.para(caption, size=9, font="Helvetica-Oblique")
        total_h = row_h * len(rows)
        self._need(total_h + 6)
        x0 = self.margin
        top = self.y
        xs = [x0]
        for w in col_widths:
            xs.append(xs[-1] + w)
        ys = [top - i * row_h for i in range(len(rows) + 1)]
        self.c.setLineWidth(0.6)
        self.c.grid(xs, ys)
        for r, row in enumerate(rows):
            self.c.setFont("Helvetica-Bold" if r == 0 else "Helvetica", size)
            for ci, cell in enumerate(row):
                self.c.drawString(xs[ci] + 3, ys[r] - row_h + 5, cell)
        self.y = ys[-1] - 10

    def image_page(self, pil_image) -> None:
        from reportlab.lib.utils import ImageReader

        self._ensure_page()
        self.c.drawImage(
            ImageReader(pil_image), 0, 0, width=self.width, height=self.height
        )
        self.y = self.margin  # page is full

    def save(self) -> None:
        self._ensure_page()
        self.c.save()


def _minimal_pdf(page_count: int, text: str) -> bytes:
    """Hand-built, valid PDF with ``page_count`` pages sharing one content stream.

    Used for the page-limit fixture: 2001 reportlab pages would exceed the
    300 KB commit budget; sharing one content stream keeps it ~180 KB.
    """
    objs: list[bytes] = []

    def add(obj: bytes) -> int:
        objs.append(obj)
        return len(objs)

    add(b"<< /Type /Catalog /Pages 2 0 R >>")  # 1
    add(b"")  # 2 placeholder for /Pages
    stream = f"BT /F1 12 Tf 72 760 Td ({text}) Tj ET".encode("latin-1")
    add(
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream"
    )  # 3
    add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")  # 4
    first_page = len(objs) + 1
    for _ in range(page_count):
        # /MediaBox and /Resources are inheritable -> declared once on /Pages to stay under 300 KB
        add(b"<< /Type /Page /Parent 2 0 R /Contents 3 0 R >>")
    kids = b" ".join(f"{first_page + i} 0 R".encode() for i in range(page_count))
    objs[1] = (
        b"<< /Type /Pages /Kids ["
        + kids
        + b"] /Count "
        + str(page_count).encode()
        + b" /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> >>"
    )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R "
        f"/ID [<{'0' * 32}> <{'0' * 32}>] >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


# --------------------------------------------------------------------------- OOXML helpers


def _docx_bytes(build: Callable, *, patches=None, extra=None) -> bytes:
    import docx

    d = docx.Document()
    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.created = FIXED_DT_NAIVE
    cp.modified = FIXED_DT_NAIVE
    cp.last_printed = FIXED_DT_NAIVE
    cp.revision = 1
    cp.comments = "AEO deterministic test corpus"
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return _normalise_zip(buf.getvalue(), patches=patches, extra=extra)


def _xlsx_bytes(build: Callable, *, patches=None, extra=None) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = FIXED_DT_NAIVE
    wb.properties.modified = FIXED_DT_NAIVE
    wb.properties.description = "AEO deterministic test corpus"
    build(wb)
    buf = io.BytesIO()
    wb.save(
        buf
    )  # NB: Workbook.save() overwrites properties.modified with now() -> pinned below
    all_patches = {"docProps/core.xml": _pin_core_modified}
    all_patches.update(patches or {})
    return _normalise_zip(buf.getvalue(), patches=all_patches, extra=extra)


def _pin_core_modified(xml: bytes) -> bytes:
    fixed = FIXED_DT_NAIVE.strftime("%Y-%m-%dT%H:%M:%SZ")
    s = re.sub(
        r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
        rf"\g<1>{fixed}\g<2>",
        xml.decode("utf-8"),
    )
    s = re.sub(
        r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rf"\g<1>{fixed}\g<2>", s
    )
    return s.encode("utf-8")


def _pptx_bytes(build: Callable, *, patches=None, extra=None) -> bytes:
    import pptx

    prs = pptx.Presentation()
    cp = prs.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.created = FIXED_DT_NAIVE
    cp.modified = FIXED_DT_NAIVE
    cp.revision = 1
    cp.comments = "AEO deterministic test corpus"
    build(prs)
    buf = io.BytesIO()
    prs.save(buf)
    return _normalise_zip(buf.getvalue(), patches=patches, extra=extra)


# --------------------------------------------------------------------------- registry


@dataclass(frozen=True)
class Fixture:
    id: str
    relpath: str  # relative to the corpus root
    build: Callable[[Path, Fixture], None]
    on_demand: bool = False  # not committed; only with --edge
    symlink_to: str | None = None  # relative target (EDGE-05)

    @property
    def is_ctx(self) -> bool:
        return self.relpath.startswith(CTX_DIRNAME + "/")


REGISTRY: list[Fixture] = []


def fixture(
    fid: str, relpath: str, *, on_demand: bool = False, symlink_to: str | None = None
):
    def deco(fn):
        REGISTRY.append(Fixture(fid, relpath, fn, on_demand, symlink_to))
        return fn

    return deco


def _write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


# =========================================================================== PDF


@fixture("PDF-01", "generated/PDF-01_single_column_headings.pdf")
def gen_pdf_01(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    w = _PdfWriter(path, "PDF-01 single column headings")
    w.heading("Ledger Service Design Notes", 20, space_before=0)
    w.para(_paragraph(rng, sentences=3))
    sections = [
        ("Overview", ["Goals", "Non-goals"]),
        ("Posting Pipeline", ["Validation", "Persistence", "Reconciliation"]),
        ("Operations", ["Monitoring", "Runbooks"]),
    ]
    for h1, subs in sections:
        w.heading(h1, 16)
        w.para(_paragraph(rng, sentences=4))
        for h2 in subs:
            w.heading(h2, 13)
            w.para(_paragraph(rng, sentences=5))
            w.para(_paragraph(rng, sentences=4))
    w.save()


@fixture("PDF-02", "generated/PDF-02_multi_column_header_footer.pdf")
def gen_pdf_02(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    w = _PdfWriter(path, "PDF-02 multi-column with repeated header and footer")
    w.header = "ACME Corp - Internal Specification v2.1 - Distribution restricted"
    w.footer = "Confidential - Page {page} of {pages}"
    w.pages_total = 6
    col_w = (w.width - 2 * w.margin - 24) / 2
    for _page in range(6):
        w._ensure_page()
        top = w.y
        for col in range(2):
            w.y = top
            x = w.margin + col * (col_w + 24)
            while w.y - 40 > w.margin:
                text = _paragraph(rng, sentences=3)
                lines = _wrap(text, 52)
                needed = len(lines) * 11.7 + 7
                if w.y - needed < w.margin:
                    break
                for line in lines:
                    w.c.setFont("Helvetica", 9)
                    w.c.drawString(x, w.y - 9, line)
                    w.y -= 11.7
                w.y -= 7
        if _page < 5:
            w.new_page()
    w.save()


@fixture("PDF-03", "generated/PDF-03_tables_and_prose.pdf")
def gen_pdf_03(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    w = _PdfWriter(path, "PDF-03 three tables with prose")
    w.heading("Capacity Plan", 18, space_before=0)
    w.para(_paragraph(rng, sentences=2))
    t1 = [["Service", "Region", "Replicas", "CPU (m)"]] + [
        [
            f"svc-{i:02d}",
            rng.choice(["eu-west-1", "us-east-1", "ap-south-1"]),
            str(rng.randint(2, 12)),
            str(rng.randint(250, 4000)),
        ]
        for i in range(1, 9)
    ]
    w.table(t1, [120, 110, 80, 80], caption="Table 1 - Service placement")
    w.para(_paragraph(rng, sentences=2))
    t2 = [["Queue", "Depth p95", "Consumers"]] + [
        [
            f"q.{rng.choice(['ingest', 'context', 'notify'])}.{i}",
            str(rng.randint(10, 900)),
            str(rng.randint(1, 8)),
        ]
        for i in range(1, 10)
    ]
    w.table(t2, [160, 100, 100], caption="Table 2 - Queue depth")
    w.new_page()
    w.para(_paragraph(rng, sentences=2))
    t3 = [["Month", "Storage GB", "Egress GB", "Cost USD", "Delta"]] + [
        [
            f"2026-{m:02d}",
            str(rng.randint(400, 900)),
            str(rng.randint(50, 300)),
            f"{rng.randint(1000, 5000)}.{rng.randint(0, 99):02d}",
            f"{rng.choice(['+', '-'])}{rng.randint(0, 15)}%",
        ]
        for m in range(1, 8)
    ]
    w.table(t3, [70, 90, 80, 90, 60], caption="Table 3 - Monthly cost")
    w.save()


VAULT_FIRST_ID = 100
VAULT_COUNT = 79
_VAULT_CHAPTERS = [
    "Introduction",
    "Account Management",
    "Eligibility Rules",
    "Transactions",
    "Posting",
    "Reporting",
    "Security",
    "Audit and Logging",
    "Operations",
]


@fixture("PDF-04", "generated/PDF-04_numbered_requirements_79p.pdf")
def gen_pdf_04(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    w = _PdfWriter(path, "VAULT Technical System Requirements")
    w.header = "VAULT Technical System Requirements - Rev 4"
    w.footer = "VAULT-TSR - Page {page} of {pages}"
    w.pages_total = VAULT_COUNT
    per_chapter = [9, 9, 9, 9, 9, 9, 9, 8, 8]  # = 79 sections
    assert sum(per_chapter) == VAULT_COUNT
    rid = VAULT_FIRST_ID
    page = 0
    for ci, (chapter, n_sec) in enumerate(
        zip(_VAULT_CHAPTERS, per_chapter, strict=True), start=1
    ):
        for si in range(1, n_sec + 1):
            if page > 0:
                w.new_page()
            page += 1
            if si == 1:
                if ci == 1:
                    w.heading("VAULT Technical System Requirements", 20, space_before=0)
                    w.para("Revision 4. " + _paragraph(rng, sentences=2))
                w.heading(f"{ci}. {chapter}", 16)
                w.para(_paragraph(rng, sentences=2))
            w.heading(f"{ci}.{si} {chapter} rule {si}", 12)
            w.para(
                f"VAULT-TSR-{rid}: The system shall "
                + _sentence(rng, _GENERAL_WORDS, 12).lower()
            )
            w.para(_paragraph(rng, sentences=3))
            w.para("Rationale: " + _paragraph(rng, sentences=2))
            rid += 1
    assert page == VAULT_COUNT
    w.save()


@fixture("PDF-05", "generated/PDF-05_scanned_image_only.pdf")
def gen_pdf_05(path: Path, fx: Fixture) -> None:
    from PIL import Image, ImageDraw

    rng = _rng(fx.id)
    w = _PdfWriter(path, "PDF-05 scanned image-only pages")
    for p in range(3):
        img = Image.new("L", (420, 594), 255)
        d = ImageDraw.Draw(img)
        y = 40
        while y < 560:
            x = 36
            while x < 380:
                seg = rng.randint(12, 60)
                d.rectangle(
                    [x, y, min(x + seg, 384), y + 5], fill=rng.choice([0, 40, 90])
                )
                x += seg + rng.randint(4, 10)
            y += rng.choice([12, 12, 12, 22])
        w.image_page(img)
        if p < 2:
            w.new_page()
    w.save()


@fixture("PDF-06", "generated/PDF-06_encrypted.pdf")
def gen_pdf_06(path: Path, fx: Fixture) -> None:
    from reportlab.lib.pdfencrypt import StandardEncryption

    rng = _rng(fx.id)
    enc = StandardEncryption(
        userPassword="evoq-user-secret",
        ownerPassword="evoq-owner-secret",
        canPrint=1,
        canModify=0,
        canCopy=0,
        canAnnotate=0,
        strength=128,
    )
    w = _PdfWriter(path, "PDF-06 encrypted", encrypt=enc)
    w.heading("Encrypted specification", 18, space_before=0)
    w.para(_paragraph(rng, sentences=4))
    w.para(
        "SEC-REQ-001: this requirement must never be extracted without the user password."
    )
    w.save()


@fixture("PDF-07", "generated/PDF-07_corrupt_truncated.pdf")
def gen_pdf_07(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    body = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
    body += rng.randbytes(6144)  # no xref, no trailer, no page tree -> corrupt
    body += b"\n%%EOF\n"
    _write(path, body)


@fixture("PDF-08", "generated/PDF-08_page_limit_2001.pdf")
def gen_pdf_08(path: Path, fx: Fixture) -> None:
    _write(
        path,
        _minimal_pdf(2001, "Page limit fixture: 2001 pages share this content stream."),
    )


_CJK_POOL = (
    "系统应当记录每一笔交易并在结算前完成校验账户余额不足时拒绝过账审计日志必须不可篡改"
    "操作员可以查询历史版本策略引擎按照优先级顺序评估规则租户之间的数据严格隔离"
    "报表在每日凌晨生成并归档密钥每九十天轮换一次会话令牌过期后需要重新认证"
)


@fixture("PDF-09", "generated/PDF-09_cjk_no_whitespace.pdf")
def gen_pdf_09(path: Path, fx: Fixture) -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    rng = _rng(fx.id)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    w = _PdfWriter(path, "PDF-09 CJK without whitespace")
    pool = _CJK_POOL
    for p in range(3):
        w._ensure_page()
        for _line in range(46):
            line = "".join(rng.choice(pool) for _ in range(44))
            w._need(13)
            w.c.setFont("STSong-Light", 10.5)
            w.c.drawString(w.margin, w.y - 10.5, line)
            w.y -= 13.5
        if p < 2:
            w.new_page()
    w.save()


# =========================================================================== DOCX


@fixture("DOCX-01", "generated/DOCX-01_headings_lists.docx")
def gen_docx_01(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)

    def build(d):
        d.add_heading("Payments Platform Handbook", 0)
        d.add_paragraph(_paragraph(rng, sentences=3))
        d.add_heading("1 Onboarding", 1)
        d.add_paragraph(_paragraph(rng, sentences=4))
        d.add_heading("1.1 Accounts", 2)
        for item in [
            "Create the tenant",
            "Assign an operator role",
            "Verify the callback URL",
        ]:
            d.add_paragraph(item, style="List Bullet")
        d.add_heading("1.2 Credentials", 2)
        for item in [
            "Request a client id",
            "Store the secret in the vault",
            "Rotate every ninety days",
        ]:
            d.add_paragraph(item, style="List Number")
        d.add_heading("2 Daily Operations", 1)
        d.add_paragraph(_paragraph(rng, sentences=4))
        d.add_heading("2.1 Posting window", 2)
        d.add_paragraph(_paragraph(rng, sentences=3))
        d.add_heading("2.1.1 Cut-off", 3)
        d.add_paragraph(_paragraph(rng, sentences=2))
        d.add_heading("2.2 Reconciliation", 2)
        d.add_paragraph(_paragraph(rng, sentences=3))
        d.add_heading("3 Incident Response", 1)
        d.add_paragraph(_paragraph(rng, sentences=4))
        d.add_heading("3.1 Severity levels", 2)
        for item in [
            "Sev-1: customer impact",
            "Sev-2: degraded service",
            "Sev-3: cosmetic",
        ]:
            d.add_paragraph(item, style="List Bullet")

    _write(path, _docx_bytes(build))


@fixture("DOCX-02", "generated/DOCX-02_merged_cells.docx")
def gen_docx_02(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)

    def build(d):
        d.add_heading("Regional totals", 1)
        d.add_paragraph(_paragraph(rng, sentences=2))
        t = d.add_table(rows=9, cols=4)
        t.style = "Table Grid"
        top = t.cell(0, 0).merge(
            t.cell(0, 3)
        )  # gridSpan=4 -> colspan; text must appear once
        top.text = "Quarterly totals"
        for ci, h in enumerate(["Region", "Q1", "Q2", "Q3"]):
            t.cell(1, ci).text = h
        left = t.cell(2, 0).merge(t.cell(3, 0))  # vMerge -> rowspan=2
        left.text = "Region A"
        for r, name in zip(
            range(4, 9),
            ["Region B", "Region C", "Region D", "Region E", "Region F"],
            strict=True,
        ):
            t.cell(r, 0).text = name
        for r in range(2, 9):
            for ci in range(1, 4):
                t.cell(r, ci).text = str(rng.randint(100, 999))
        d.add_paragraph("Totals exclude reversals.")

    _write(path, _docx_bytes(build))


_FOOTNOTES_XML = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    b'<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    b'<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    b'<w:footnote w:id="1"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
    b'<w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> Footnote FN-ALPHA: settlement windows are defined by the clearing house.</w:t></w:r></w:p></w:footnote>'
    b"</w:footnotes>"
)

_TEXTBOX_RUN = (
    '<w:r><w:pict xmlns:v="urn:schemas-microsoft-com:vml"><v:shape id="tb1" type="#_x0000_t202" '
    'style="position:absolute;width:200pt;height:40pt"><v:textbox><w:txbxContent><w:p><w:r>'
    "<w:t>TEXTBOX-BRAVO: callout shown once in the extracted body.</w:t></w:r></w:p></w:txbxContent>"
    "</v:textbox></v:shape></w:pict></w:r>"
)


@fixture("DOCX-03", "generated/DOCX-03_header_footer_footnote_textbox.docx")
def gen_docx_03(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)

    def build(d):
        sec = d.sections[0]
        sec.header.paragraphs[0].text = "HEADER-CHARLIE: Settlement Operations Manual"
        sec.footer.paragraphs[0].text = "FOOTER-DELTA: Internal use only"
        d.add_heading("Settlement windows", 1)
        d.add_paragraph(_paragraph(rng, sentences=3))
        d.add_paragraph("Windows close at 17:00 local time.@@FOOTNOTE@@")
        d.add_paragraph("@@TEXTBOX@@")
        d.add_heading("Exceptions", 2)
        d.add_paragraph(_paragraph(rng, sentences=3))

    def patch_document(xml: bytes) -> bytes:
        s = xml.decode("utf-8")
        s = _sub_once(
            r"@@FOOTNOTE@@</w:t></w:r>",
            '</w:t></w:r><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
            '<w:footnoteReference w:id="1"/></w:r>',
            s,
        )
        s = _sub_once(r"<w:r><w:t>@@TEXTBOX@@</w:t></w:r>", _TEXTBOX_RUN, s)
        return s.encode("utf-8")

    def patch_rels(xml: bytes) -> bytes:
        s = xml.decode("utf-8")
        rel = (
            '<Relationship Id="rIdFootnotes1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
            'Target="footnotes.xml"/></Relationships>'
        )
        return _sub_once(r"</Relationships>", rel, s).encode("utf-8")

    def patch_ct(xml: bytes) -> bytes:
        s = xml.decode("utf-8")
        ov = (
            '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.wordprocessingml.footnotes+xml"/></Types>'
        )
        return _sub_once(r"</Types>", ov, s).encode("utf-8")

    _write(
        path,
        _docx_bytes(
            build,
            patches={
                "word/document.xml": patch_document,
                "word/_rels/document.xml.rels": patch_rels,
                "[Content_Types].xml": patch_ct,
            },
            extra={"word/footnotes.xml": _FOOTNOTES_XML},
        ),
    )


@fixture("DOCX-04", "generated/DOCX-04_tracked_changes.docx")
def gen_docx_04(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)

    def build(d):
        d.add_heading("Response time requirement", 1)
        d.add_paragraph(_paragraph(rng, sentences=2))
        d.add_paragraph(
            "The service shall respond within @@TRACKED@@ under nominal load."
        )
        d.add_paragraph(_paragraph(rng, sentences=2))

    def patch_document(xml: bytes) -> bytes:
        s = xml.decode("utf-8")
        replacement = (
            "The service shall respond within </w:t></w:r>"
            '<w:del w:id="1" w:author="Reviewer" w:date="2026-09-11T00:00:00Z">'
            "<w:r><w:delText>DELETED-ECHO 500 ms</w:delText></w:r></w:del>"
            '<w:ins w:id="2" w:author="Reviewer" w:date="2026-09-11T00:00:00Z">'
            "<w:r><w:t>INSERTED-FOXTROT 200 ms</w:t></w:r></w:ins>"
            '<w:r><w:t xml:space="preserve"> under nominal load.'
        )
        return _sub_once(
            r"The service shall respond within @@TRACKED@@ under nominal load\.",
            replacement,
            s,
        ).encode()

    _write(path, _docx_bytes(build, patches={"word/document.xml": patch_document}))


@fixture("DOCX-05", "generated/DOCX-05_numbered_headings_inferred.docx")
def gen_docx_05(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    outline = [
        ("1. Introduction", 3),
        ("1.1 Scope", 2),
        ("1.2 Definitions", 2),
        ("2. Functional Requirements", 2),
        ("2.1 Login", 3),
        ("2.2 Session handling", 3),
        ("2.3 Password reset", 2),
        ("3. Non-functional Requirements", 2),
        ("3.1 Performance", 3),
        ("3.2 Availability", 2),
    ]

    def build(d):
        d.add_paragraph("Customer Portal Requirements Specification")
        d.add_paragraph("Document id PORTAL-SRS-7, revision 2.")
        n = 1
        for title, paras in outline:
            d.add_paragraph(
                title
            )  # Normal style, numbered text -> heading must be inferred
            for _ in range(paras):
                if title[0].isdigit() and "." in title[:4] and rng.random() < 0.5:
                    d.add_paragraph(
                        f"REQ-{n:03d}: The portal shall "
                        + _sentence(rng, _GENERAL_WORDS, 10).lower()
                    )
                    n += 1
                else:
                    d.add_paragraph(_paragraph(rng, sentences=3))

    _write(path, _docx_bytes(build))


@fixture("DOCX-06", "generated/DOCX-06_zip_renamed.docx")
def gen_docx_06(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    payload = {
        "README.txt": b"This is a plain zip archive renamed to .docx. It is not a Word document.\n",
        "data/notes.txt": _paragraph(rng, sentences=3).encode(),
    }
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in payload.items():
            zi = zipfile.ZipInfo(name, date_time=ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 3
            zi.external_attr = 0o644 << 16
            z.writestr(zi, data, compresslevel=ZIP_LEVEL)
    _write(path, buf.getvalue())


# =========================================================================== XLSX


@fixture("XLSX-01", "generated/XLSX-01_three_sheets.xlsx")
def gen_xlsx_01(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)

    def build(wb):
        ws = wb.active
        ws.title = "Requirements"
        ws.append(["ID", "Title", "Priority", "Owner"])
        for i in range(1, 21):
            ws.append(
                [
                    f"REQ-{i:03d}",
                    _sentence(rng, _GENERAL_WORDS, 6),
                    rng.choice(["High", "Medium", "Low"]),
                    rng.choice(["alice", "bob", "chen", "dara"]),
                ]
            )
        ws2 = wb.create_sheet("Risks")
        ws2.append(["ID", "Description", "Likelihood", "Impact"])
        for i in range(1, 13):
            ws2.append(
                [
                    f"RISK-{i:02d}",
                    _sentence(rng, _GENERAL_WORDS, 8),
                    rng.randint(1, 5),
                    rng.randint(1, 5),
                ]
            )
        ws3 = wb.create_sheet("Glossary")
        ws3.append(["Term", "Definition"])
        for term in [
            "Ledger",
            "Posting",
            "Reversal",
            "Tenant",
            "Operator",
            "Snapshot",
            "Manifest",
            "Chunk",
            "Vector",
            "Requirement",
            "Audit",
            "Policy",
            "Session",
            "Token",
            "Queue",
        ]:
            ws3.append([term, _sentence(rng, _GENERAL_WORDS, 9)])

    _write(path, _xlsx_bytes(build))


@fixture("XLSX-02", "generated/XLSX-02_formulas_no_cache.xlsx")
def gen_xlsx_02(path: Path, fx: Fixture) -> None:
    def build(wb):
        ws = wb.active
        ws.title = "Budget"
        ws.append(["Item", "Amount", "Adjusted"])
        for item, amount in [
            ("Compute", 400),
            ("Storage", 250),
            ("Egress", 150),
            ("Support", 200),
        ]:
            ws.append([item, amount, None])
        for r in range(2, 6):
            ws.cell(row=r, column=3, value=f"=B{r}*1.1")
        ws.cell(row=6, column=1, value="Total")
        ws.cell(
            row=6, column=2, value="=SUM(B2:B5)"
        )  # gets a cached <v> via the zip patch below
        ws.cell(row=7, column=1, value="Average")
        ws.cell(
            row=7, column=2, value="=AVERAGE(B2:B5)"
        )  # stays without a cached value

    def patch_sheet(xml: bytes) -> bytes:
        s = xml.decode("utf-8")
        # openpyxl writes `<c r="B6"><f>SUM(B2:B5)</f><v></v></c>`; give exactly that one cell a cached value.
        return _sub_once(
            r'(<c r="B6"[^>]*><f>SUM\(B2:B5\)</f>)<v\s*/?>(?:</v>)?',
            r"\g<1><v>1000</v>",
            s,
        ).encode()

    _write(path, _xlsx_bytes(build, patches={"xl/worksheets/sheet1.xml": patch_sheet}))


@fixture("XLSX-03", "generated/XLSX-03_merged_header_empty_rows_cols.xlsx")
def gen_xlsx_03(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)

    def build(wb):
        ws = wb.active
        ws.title = "Report"
        ws["A1"] = "Quarterly report - FY26"
        ws.merge_cells("A1:D1")
        # row 2 intentionally empty; column C intentionally empty
        ws.append([])
        ws.append(["Region", "Q1", None, "Q3"])
        for region in ["North", "South", "East", "West", "Central", "Islands"]:
            ws.append([region, rng.randint(100, 999), None, rng.randint(100, 999)])
        ws.append([])
        ws.append([])
        for region in ["Overseas", "Partners", "Other"]:
            ws.append([region, rng.randint(10, 99), None, rng.randint(10, 99)])
        hidden = wb.create_sheet("Scratch")
        hidden.sheet_state = "hidden"
        hidden.append(["scratch", "HIDDEN-GOLF value"])

    _write(path, _xlsx_bytes(build))


XLSX04_ROWS = 200_001


@fixture("XLSX-04", "generated/XLSX-04_rows_200001.xlsx")
def gen_xlsx_04(path: Path, fx: Fixture) -> None:
    def build(wb):
        ws = wb.active
        ws.title = "Data"
        ws.append(["id", "value"])  # replaced wholesale by the patch below

    def sheet_xml(_: bytes) -> bytes:
        head = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<dimension ref="A1:B{XLSX04_ROWS + 1}"/><sheetData>'
            '<row r="1"><c r="A1" t="inlineStr"><is><t>id</t></is></c>'
            '<c r="B1" t="inlineStr"><is><t>value</t></is></c></row>'
        )
        # rows without r= attributes are valid SpreadsheetML and compress to almost nothing
        row = "<row><c><v>1</v></c><c><v>2</v></c></row>"
        tail = "</sheetData></worksheet>"
        return (head + row * XLSX04_ROWS + tail).encode()

    _write(path, _xlsx_bytes(build, patches={"xl/worksheets/sheet1.xml": sheet_xml}))


@fixture("XLSX-05", "generated/XLSX-05_dates_percent_iso.xlsx")
def gen_xlsx_05(path: Path, fx: Fixture) -> None:
    def build(wb):
        ws = wb.active
        ws.title = "Schedule"
        ws.append(["Milestone", "Date", "Timestamp", "Done", "Progress", "Cost"])
        rows = [
            (
                "Kick-off",
                date(2026, 3, 15),
                datetime(2026, 3, 15, 9, 30),
                True,
                0.125,
                1200.5,
            ),
            (
                "Design freeze",
                date(2026, 5, 1),
                datetime(2026, 5, 1, 17, 0),
                False,
                0.5,
                3400.0,
            ),
            (
                "Go-live",
                date(2026, 9, 30),
                datetime(2026, 9, 30, 6, 45, 30),
                False,
                0.875,
                99999.99,
            ),
        ]
        for r in rows:
            ws.append(list(r))
        for r in range(2, 5):
            ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
            ws.cell(row=r, column=3).number_format = "yyyy-mm-dd hh:mm:ss"
            ws.cell(row=r, column=5).number_format = "0.00%"
            ws.cell(row=r, column=6).number_format = "#,##0.00"

    _write(path, _xlsx_bytes(build))


# =========================================================================== CSV


@fixture("CSV-01", "generated/CSV-01_bom_utf8.csv")
def gen_csv_01(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    lines = ["id,title,owner,status"]
    for i in range(1, 13):
        lines.append(
            f"{i},{_sentence(rng, _GENERAL_WORDS, 5)[:-1]},{rng.choice(['alice', 'bob', 'chen'])},"
            f"{rng.choice(['open', 'done', 'blocked'])}"
        )
    _write(path, ("\n".join(lines) + "\n").encode("utf-8-sig"))


@fixture("CSV-02", "generated/CSV-02_semicolon_quoted_newlines.csv")
def gen_csv_02(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    lines = ["key;description;notes"]
    lines.append('CFG-1;"Retry policy; exponential backoff";"line one\nline two"')
    lines.append('CFG-2;"Timeout in seconds";"single line"')
    lines.append('CFG-3;"Quoted ""inner"" text";"first\nsecond\nthird"')
    for i in range(4, 10):
        lines.append(f"CFG-{i};{_sentence(rng, _GENERAL_WORDS, 6)[:-1]};plain")
    _write(path, ("\n".join(lines) + "\n").encode("utf-8"))


@fixture("CSV-03", "generated/CSV-03_utf16.csv")
def gen_csv_03(path: Path, fx: Fixture) -> None:
    lines = [
        "id,name,city",
        "1,Señora Muñoz,Málaga",
        "2,Jörg Bäcker,Köln",
        "3,山田 太郎,東京",
        "4,Zoë Ó Briain,Baile Átha Cliath",
    ]
    _write(
        path, ("\n".join(lines) + "\n").encode("utf-16")
    )  # BOM + platform-native (LE) order


@fixture("CSV-04", "generated/CSV-04_ragged_rows.csv")
def gen_csv_04(path: Path, fx: Fixture) -> None:
    lines = ["a,b,c,d", "1,2,3", "4,5,6,7", "8,9,10,11,12", "13,14", "15,16,17,18"]
    _write(path, ("\n".join(lines) + "\n").encode("utf-8"))


# =========================================================================== PPTX


@fixture("PPTX-01", "generated/PPTX-01_six_slides_notes.pptx")
def gen_pptx_01(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)

    def build(prs):
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = "Context Engine Roadmap"
        s.placeholders[1].text = "Quarterly review"
        s.notes_slide.notes_text_frame.text = "Speaker note 1: welcome and agenda."
        topics = ["Chunking", "Embedding", "Retrieval", "Synthesis", "Operations"]
        for i, topic in enumerate(topics, start=2):
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = topic
            tf = s.placeholders[1].text_frame
            tf.text = _sentence(rng, _GENERAL_WORDS, 6)
            for _ in range(rng.randint(2, 3)):
                p = tf.add_paragraph()
                p.text = _sentence(rng, _GENERAL_WORDS, 6)
                p.level = 1
            s.notes_slide.notes_text_frame.text = f"Speaker note {i}: " + _sentence(
                rng, _GENERAL_WORDS, 8
            )

    _write(path, _pptx_bytes(build))


@fixture("PPTX-02", "generated/PPTX-02_table_grouped_shapes_hidden.pptx")
def gen_pptx_02(path: Path, fx: Fixture) -> None:
    from pptx.util import Inches, Pt

    rng = _rng(fx.id)

    def build(prs):
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = "Service inventory"
        s.placeholders[1].text = "Tables, groups and a hidden slide"
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "Inventory table"
        tbl = s.shapes.add_table(
            4, 3, Inches(0.5), Inches(1.5), Inches(9), Inches(2)
        ).table
        for ci, h in enumerate(["Service", "Owner", "Tier"]):
            tbl.cell(0, ci).text = h
        for r, svc in enumerate(["ingestion", "context", "identity"], start=1):
            tbl.cell(r, 0).text = svc
            tbl.cell(r, 1).text = rng.choice(["alice", "bob", "chen"])
            tbl.cell(r, 2).text = rng.choice(["gold", "silver"])
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "Grouped shapes"
        grp = s.shapes.add_group_shape()
        tb1 = grp.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        tb1.text_frame.text = "GROUP-HOTEL left box text"
        tb2 = grp.shapes.add_textbox(Inches(5), Inches(2), Inches(4), Inches(1))
        tb2.text_frame.text = "GROUP-INDIA right box text"
        tb2.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "Hidden slide"
        box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
        box.text_frame.text = "HIDDEN-JULIET content of a hidden slide"
        s._element.set("show", "0")

    _write(path, _pptx_bytes(build))


@fixture("PPTX-03", "generated/PPTX-03_slide_limit_101.pptx")
def gen_pptx_03(path: Path, fx: Fixture) -> None:
    def build(prs):
        for i in range(1, 102):
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = f"Slide {i}"

    _write(path, _pptx_bytes(build))


# =========================================================================== MD / TXT

_MD_01 = """# Ingestion Engine Guide

The engine turns uploads into canonical Markdown. SPEC-001 requires every body to carry structure markers.

## Supported formats

| Format | Engine | Markers |
|---|---|---|
| PDF | pdfplumber | `--- Page N ---` |
| XLSX | openpyxl | `--- Sheet: name ---` |
| PPTX | python-pptx | `--- Slide N ---` |

## Configuration

Settings are read from the environment. SPEC-002 requires defaults to live in one module.

```python
from evoq_ingestion import Limits

limits = Limits(max_pdf_pages=2000)
```

### Environment table

| Variable | Default |
|---|---|
| `PDF_ENGINE` | `pdfplumber` |
| `XLSX_MAX_ROWS_PER_SHEET` | `200000` |

## Failure codes

Every refusal is a lower_snake code such as `encrypted_pdf` or `file_too_large`.

```json
{"status": "FAILED", "failure_code": "encrypted_pdf"}
```

## Links

See the [contract](../canonical_document.md) and the [markers grammar](markers.md).
"""

_MD_02 = """Setext Heading Level One
========================

Intro paragraph under a setext H1. Nothing here uses ATX syntax.

Second Level Heading
--------------------

- Top-level bullet one
  - Nested bullet one-a
    - Deep bullet one-a-i
    - Deep bullet one-a-ii
  - Nested bullet one-b
- Top-level bullet two
  1. Ordered child two-a
  2. Ordered child two-b
     - Mixed nested bullet under ordered
- Top-level bullet three

Another Second Level
--------------------

> A block quote spanning
> two lines.

Closing paragraph with *emphasis* and **strong** text.
"""

_CTX_01 = """# Tiny document

This document has one heading and two sentences. It fits in a single child chunk.
"""


@fixture("MD-01", "generated/MD-01_atx_fences_tables.md")
def gen_md_01(path: Path, fx: Fixture) -> None:
    _write(path, _MD_01.encode())


@fixture("MD-02", "generated/MD-02_setext_nested_lists.md")
def gen_md_02(path: Path, fx: Fixture) -> None:
    _write(path, _MD_02.encode())


@fixture("TXT-01", "generated/TXT-01_two_topics_paragraphs.txt")
def gen_txt_01(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    paras = [_paragraph(rng, _DATABASE_WORDS, sentences=5) for _ in range(6)]
    paras += [_paragraph(rng, _GARDEN_WORDS, sentences=5) for _ in range(6)]
    _write(path, ("\n\n".join(paras) + "\n").encode())


@fixture("TXT-02", "generated/TXT-02_utf16.txt")
def gen_txt_02(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    paras = [
        "UTF-16 sample: Zoë, Jörg, Señora Muñoz, 東京, Ærø.",
        _paragraph(rng, sentences=3),
        _paragraph(rng, sentences=3),
        "Ende des Dokuments — fin du document.",
    ]
    _write(path, ("\n\n".join(paras) + "\n").encode("utf-16"))


@fixture("TXT-03", "generated/TXT-03_binary_content.txt")
def gen_txt_03(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    _write(path, rng.randbytes(4096))


@fixture("TXT-04", "generated/TXT-04_empty.txt")
def gen_txt_04(path: Path, fx: Fixture) -> None:
    _write(path, b"")


@fixture("TXT-05", "generated/TXT-05_base64_blob.txt")
def gen_txt_05(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    _write(
        path, base64.b64encode(rng.randbytes(30_000)) + b"\n"
    )  # one 40,000-char line, no whitespace


# =========================================================================== HTML (Confluence)

_HTML_01 = """<h1>Overview</h1>
<p>The ingestion job renders Confluence storage-format pages into canonical Markdown.</p>
<ac:structured-macro ac:name="toc" ac:schema-version="1"><ac:parameter ac:name="maxLevel">3</ac:parameter></ac:structured-macro>
<h2>Getting started</h2>
<p>Install the package and run the extractor once against a sample page. Keep <strong>secrets</strong> out of the page body.</p>
<ac:structured-macro ac:name="code" ac:schema-version="1">
  <ac:parameter ac:name="language">python</ac:parameter>
  <ac:parameter ac:name="title">handler.py</ac:parameter>
  <ac:plain-text-body><![CDATA[def handler(event):
    return {"ok": True, "pipes": "a | b"}]]></ac:plain-text-body>
</ac:structured-macro>
<h2>Configuration</h2>
<ul>
  <li>Set <code>CONFLUENCE_BODY_FORMAT</code> to <em>storage</em> for CDATA-safe rendering.</li>
  <li>Set <code>CONFLUENCE_PAGE_SIZE</code> to 25.</li>
</ul>
<h3>Environment</h3>
<ac:structured-macro ac:name="info" ac:schema-version="1">
  <ac:rich-text-body><p>INFO-KILO: the info panel body must be rendered.</p></ac:rich-text-body>
</ac:structured-macro>
<h3>Status</h3>
<p>Current status: <ac:structured-macro ac:name="status" ac:schema-version="1"><ac:parameter ac:name="colour">Green</ac:parameter><ac:parameter ac:name="title">APPROVED</ac:parameter></ac:structured-macro></p>
"""

_HTML_02 = """<h1>Component matrix</h1>
<p>The outer table has one cell that contains an inner table.</p>
<table>
  <tbody>
    <tr><th>Component</th><th>Details</th></tr>
    <tr>
      <td>Authentication</td>
      <td>
        <table>
          <tbody>
            <tr><th>Key</th><th>Value</th></tr>
            <tr><td>Protocol</td><td>OIDC</td></tr>
            <tr><td>Token TTL</td><td>15 minutes</td></tr>
          </tbody>
        </table>
      </td>
    </tr>
    <tr><td>Storage</td><td>Object store with versioning; NESTED-LIMA marker cell.</td></tr>
  </tbody>
</table>
<p>Closing paragraph after the table.</p>
"""

_HTML_03 = """<h1>Release schedule</h1>
<table>
  <tbody>
    <tr><th>Quarter</th><th colspan="2">Deliverables</th><th>Owner</th></tr>
    <tr><td rowspan="2">Q1</td><td>Chunking rewrite</td><td>SPAN-MIKE</td><td>alice</td></tr>
    <tr><td>Parent-child</td><td>topic refiner</td><td>bob</td></tr>
    <tr><td>Q2</td><td colspan="2">Observability alignment</td><td>chen</td></tr>
  </tbody>
</table>
"""

_HTML_04 = """<div id="main-content" class="wiki-content">
<h1>Architecture overview</h1>
<p>See the <a href="https://example.atlassian.net/wiki/spaces/ENG/pages/123456/Design+notes">Design notes</a> and the <a href="https://example.atlassian.net/browse/VAULT-42">tracking issue</a>.</p>
<p><img class="confluence-embedded-image" src="/wiki/download/attachments/123456/diagram.png?version=2" data-linked-resource-id="789" data-linked-resource-type="attachment" alt="Architecture diagram" /></p>
<h2>Deployment</h2>
<p>Every service ships as one container.</p>
<p><img class="confluence-embedded-image" src="/wiki/download/attachments/123456/deploy.png" data-linked-resource-id="790" data-linked-resource-type="attachment" alt="Deployment topology" /></p>
</div>
"""


@fixture("HTML-01", "generated/HTML-01_storage_headings_code_macro.html")
def gen_html_01(path: Path, fx: Fixture) -> None:
    _write(path, _HTML_01.encode())


@fixture("HTML-02", "generated/HTML-02_nested_tables.html")
def gen_html_02(path: Path, fx: Fixture) -> None:
    _write(path, _HTML_02.encode())


@fixture("HTML-03", "generated/HTML-03_rowspan_colspan.html")
def gen_html_03(path: Path, fx: Fixture) -> None:
    _write(path, _HTML_03.encode())


@fixture("HTML-04", "generated/HTML-04_export_view_images_links.html")
def gen_html_04(path: Path, fx: Fixture) -> None:
    _write(path, _HTML_04.encode())


# =========================================================================== ADF (Jira issue payloads)


def _adf_text(text: str, *marks: dict) -> dict:
    node = {"type": "text", "text": text}
    if marks:
        node["marks"] = list(marks)
    return node


def _adf_para(*content) -> dict:
    return {"type": "paragraph", "content": list(content)}


def _adf_heading(level: int, text: str) -> dict:
    return {"type": "heading", "attrs": {"level": level}, "content": [_adf_text(text)]}


def _adf_doc(*content) -> dict:
    return {"version": 1, "type": "doc", "content": list(content)}


def _jira_issue(
    key: str, summary: str, description: dict, comments: list[dict], *, issue_id: str
) -> dict:
    return {
        "id": issue_id,
        "key": key,
        "self": f"https://example.atlassian.net/rest/api/3/issue/{issue_id}",
        "fields": {
            "summary": summary,
            "issuetype": {"name": "Story"},
            "status": {"name": "In Progress"},
            "priority": {"name": "High"},
            "project": {"key": key.split("-")[0], "name": "Vault"},
            "labels": ["adf-corpus"],
            "created": "2026-09-01T09:00:00.000+0000",
            "updated": "2026-09-11T00:00:00.000+0000",
            "reporter": {"displayName": "Corpus Reporter", "accountId": "acc-reporter"},
            "assignee": None,
            "description": description,
            "comment": {
                "comments": comments,
                "total": len(comments),
                "maxResults": len(comments),
                "startAt": 0,
            },
        },
    }


def _adf_comment(i: int, author: str, body: dict) -> dict:
    return {
        "id": str(20000 + i),
        "author": {"displayName": author, "accountId": f"acc-{author.lower()}"},
        "created": f"2026-09-{i + 1:02d}T10:00:00.000+0000",
        "updated": f"2026-09-{i + 1:02d}T10:00:00.000+0000",
        "body": body,
    }


@fixture("ADF-01", "generated/ADF-01_description_nodes.json")
def gen_adf_01(path: Path, fx: Fixture) -> None:
    desc = _adf_doc(
        _adf_heading(2, "Goal"),
        _adf_para(
            _adf_text("The vault must rotate keys every ninety days."),
            _adf_text(" This sentence appears "),
            _adf_text("exactly once", {"type": "strong"}),
            _adf_text(" in the rendered body."),
        ),
        _adf_heading(3, "Acceptance"),
        {
            "type": "bulletList",
            "content": [
                {
                    "type": "listItem",
                    "content": [
                        _adf_para(
                            _adf_text("Rotation is logged with "),
                            _adf_text("AUDIT-ROT-1", {"type": "code"}),
                        )
                    ],
                },
                {
                    "type": "listItem",
                    "content": [
                        _adf_para(
                            _adf_text("Old keys are retired after "),
                            _adf_text("grace period", {"type": "em"}),
                        )
                    ],
                },
                {
                    "type": "listItem",
                    "content": [
                        _adf_para(
                            _adf_text("Runbook: "),
                            _adf_text(
                                "link",
                                {
                                    "type": "link",
                                    "attrs": {
                                        "href": "https://example.atlassian.net/wiki/x"
                                    },
                                },
                            ),
                        )
                    ],
                },
            ],
        },
        {
            "type": "panel",
            "attrs": {"panelType": "info"},
            "content": [
                _adf_para(_adf_text("PANEL-NOVEMBER: informational panel body."))
            ],
        },
        {
            "type": "blockquote",
            "content": [
                _adf_para(_adf_text("Quoted rationale from the security review."))
            ],
        },
        {"type": "rule"},
        _adf_para(_adf_text("Closing paragraph.")),
    )
    issue = _jira_issue(
        "VAULT-42", "Rotate vault keys automatically", desc, [], issue_id="10042"
    )
    _write(
        path,
        (
            json.dumps(issue, indent=2, ensure_ascii=False, sort_keys=True) + "\n"
        ).encode(),
    )


@fixture("ADF-02", "generated/ADF-02_five_comments.json")
def gen_adf_02(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    desc = _adf_doc(
        _adf_para(_adf_text("Short description; the discussion lives in the comments."))
    )
    authors = ["Alice", "Bob", "Chen", "Dara", "Alice"]
    comments = []
    for i, author in enumerate(authors, start=1):
        body = _adf_doc(
            _adf_para(_adf_text(f"COMMENT-{i}: " + _paragraph(rng, sentences=3))),
            _adf_para(_adf_text(_paragraph(rng, sentences=2))),
        )
        comments.append(_adf_comment(i, author, body))
    issue = _jira_issue(
        "VAULT-57",
        "Discussion thread about posting cut-off",
        desc,
        comments,
        issue_id="10057",
    )
    _write(
        path,
        (
            json.dumps(issue, indent=2, ensure_ascii=False, sort_keys=True) + "\n"
        ).encode(),
    )


@fixture("ADF-03", "generated/ADF-03_tasklist_table_code.json")
def gen_adf_03(path: Path, fx: Fixture) -> None:
    def cell(kind: str, text: str, **attrs) -> dict:
        node = {"type": kind, "content": [_adf_para(_adf_text(text))]}
        if attrs:
            node["attrs"] = attrs
        return node

    desc = _adf_doc(
        _adf_heading(2, "Checklist"),
        {
            "type": "taskList",
            "attrs": {"localId": "tl-1"},
            "content": [
                {
                    "type": "taskItem",
                    "attrs": {"localId": "ti-1", "state": "DONE"},
                    "content": [_adf_text("Write the migration")],
                },
                {
                    "type": "taskItem",
                    "attrs": {"localId": "ti-2", "state": "TODO"},
                    "content": [_adf_text("Review the migration")],
                },
                {
                    "type": "taskItem",
                    "attrs": {"localId": "ti-3", "state": "TODO"},
                    "content": [_adf_text("Deploy to staging")],
                },
            ],
        },
        _adf_heading(2, "Matrix"),
        {
            "type": "table",
            "attrs": {"isNumberColumnEnabled": False, "layout": "default"},
            "content": [
                {
                    "type": "tableRow",
                    "content": [
                        cell("tableHeader", "Env"),
                        cell("tableHeader", "Owner"),
                        cell("tableHeader", "Note"),
                    ],
                },
                {
                    "type": "tableRow",
                    "content": [
                        cell("tableCell", "staging"),
                        cell("tableCell", "alice"),
                        cell("tableCell", "TABLE-OSCAR"),
                    ],
                },
                {
                    "type": "tableRow",
                    "content": [
                        cell("tableCell", "prod"),
                        cell("tableCell", "bob"),
                        cell("tableCell", "pipe | inside"),
                    ],
                },
            ],
        },
        _adf_heading(2, "Snippet"),
        {
            "type": "codeBlock",
            "attrs": {"language": "json"},
            "content": [_adf_text('{"retries": 3, "backoff": "exponential"}')],
        },
        {
            "type": "orderedList",
            "attrs": {"order": 3},
            "content": [
                {"type": "listItem", "content": [_adf_para(_adf_text("Third step"))]},
                {"type": "listItem", "content": [_adf_para(_adf_text("Fourth step"))]},
            ],
        },
        {
            "type": "decisionList",
            "attrs": {"localId": "dl-1"},
            "content": [
                {
                    "type": "decisionItem",
                    "attrs": {"localId": "di-1", "state": "DECIDED"},
                    "content": [_adf_text("DECISION-PAPA: use exponential backoff")],
                },
            ],
        },
        {
            "type": "nestedExpand",
            "attrs": {"title": "Details"},
            "content": [
                _adf_para(_adf_text("EXPAND-QUEBEC: hidden by default in Jira."))
            ],
        },
        _adf_para(
            _adf_text("Status "),
            {
                "type": "status",
                "attrs": {"text": "APPROVED", "color": "green", "localId": "st-1"},
            },
            _adf_text(" assigned to "),
            {"type": "mention", "attrs": {"id": "acc-alice", "text": "@Alice"}},
            _adf_text(" on "),
            {"type": "date", "attrs": {"timestamp": "1789516800000"}},
            _adf_text(" "),
            {"type": "emoji", "attrs": {"shortName": ":thumbsup:", "text": "👍"}},
            {"type": "hardBreak"},
            _adf_text("after a hard break."),
        ),
        {
            "type": "mediaSingle",
            "attrs": {"layout": "center"},
            "content": [
                {
                    "type": "media",
                    "attrs": {
                        "id": "media-1",
                        "type": "file",
                        "collection": "c1",
                        "alt": "Diagram",
                    },
                }
            ],
        },
        {
            "type": "extension",
            "attrs": {
                "extensionType": "com.atlassian.confluence.macro.core",
                "extensionKey": "jira-chart",
            },
        },
    )
    issue = _jira_issue(
        "VAULT-61", "Task lists, tables and code in ADF", desc, [], issue_id="10061"
    )
    _write(
        path,
        (
            json.dumps(issue, indent=2, ensure_ascii=False, sort_keys=True) + "\n"
        ).encode(),
    )


# =========================================================================== EDGE


@fixture(
    "EDGE-01", f"generated/{EDGE_DIRNAME}/EDGE-01_sparse_201mib.pdf", on_demand=True
)
def gen_edge_01(path: Path, fx: Fixture) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "wb") as f:
        f.write(b"%PDF-1.4\n")
        f.seek(SPARSE_BYTES - 1)
        f.write(b"\0")


@fixture("EDGE-02", "generated/EDGE-02_unsupported_ext.xyz")
def gen_edge_02(path: Path, fx: Fixture) -> None:
    _write(path, b"This file has an extension outside the supported set.\n")


@fixture("EDGE-03", "generated/EDGE-03_ünïcødé_文件名.md")
def gen_edge_03(path: Path, fx: Fixture) -> None:
    _write(
        path,
        b"# Unicode filename\n\nThe file name carries Latin-1 accents and CJK characters.\n",
    )


@fixture("EDGE-04", "generated/EDGE-04_zero_byte.pdf")
def gen_edge_04(path: Path, fx: Fixture) -> None:
    _write(path, b"")


@fixture(
    "EDGE-05",
    "generated/EDGE-05_symlink.pdf",
    symlink_to="PDF-01_single_column_headings.pdf",
)
def gen_edge_05(path: Path, fx: Fixture) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_symlink() or path.exists():
        path.unlink()
    os.symlink(fx.symlink_to, path)


@fixture("EDGE-06", "generated/EDGE-06_duplicate_of_md01.md")
def gen_edge_06(path: Path, fx: Fixture) -> None:
    _write(
        path, _MD_01.encode()
    )  # byte-identical to MD-01 -> identical content_hash / chunk keys


# EDGE-07 (per-file timeout) re-uses PDF-04's file with a tiny timeout — see scenarios.yaml `engine_args`.


# =========================================================================== CTX (chunk-only inputs)


@fixture("CTX-01", "ctx/CTX-01_tiny.md")
def gen_ctx_01(path: Path, fx: Fixture) -> None:
    _write(path, _CTX_01.encode())


@fixture("CTX-02", "ctx/CTX-02_giant_section.md")
def gen_ctx_02(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    paras = [
        _paragraph(rng, sentences=6) for _ in range(40)
    ]  # ~3,000+ tokens under one heading
    body = "# 4. Transaction posting\n\n" + "\n\n".join(paras) + "\n"
    _write(path, body.encode())


@fixture("CTX-03", "ctx/CTX-03_prose_requirements_with_table.md")
def gen_ctx_03(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    lines = [
        "# System requirements",
        "",
        "Prose requirements follow; each id starts its own line.",
    ]
    for i in range(1, 13):
        lines += [
            "",
            f"SYS-REQ-{i:03d}: The system shall "
            + _sentence(rng, _GENERAL_WORDS, 11).lower()
            + " "
            + _sentence(rng, _GENERAL_WORDS, 9),
        ]
    lines += [
        "",
        "## Reference table",
        "",
        "| ID | Statement | Priority |",
        "|---|---|---|",
        "| SYS-REQ-013 | The audit log shall be immutable. | High |",
        "| - | Rows without ids must not start requirement blocks. | Low |",
        "| - | Neither shall this one. | Low |",
        "",
        "Closing note: the table row carrying SYS-REQ-013 is the only tabular requirement.",
        "",
    ]
    _write(path, "\n".join(lines).encode())


@fixture("CTX-04", "ctx/CTX-04_mixed_page_heading_markers.md")
def gen_ctx_04(path: Path, fx: Fixture) -> None:
    rng = _rng(fx.id)
    parts = [
        "--- Page 1 ---",
        "",
        "# 1. Introduction",
        "",
        _paragraph(rng, sentences=4),
        "",
        "## 1.1 Purpose",
        "",
        _paragraph(rng, sentences=4),
        "",
        "--- Page 2 ---",
        "",
        "## 1.2 Scope",
        "",
        _paragraph(rng, sentences=5),
        "",
        _paragraph(rng, sentences=3),
        "",
        "--- Page 3 ---",
        "",
        "# 2. Transactions",
        "",
        _paragraph(rng, sentences=3),
        "",
        "## 2.1 Posting",
        "",
        _paragraph(rng, sentences=5),
        "",
        "--- Page 4 ---",
        "",
        "## 2.2 Reversal",
        "",
        _paragraph(rng, sentences=4),
        "",
        "### 2.2.1 Same-day reversal",
        "",
        _paragraph(rng, sentences=3),
        "",
    ]
    _write(path, "\n".join(parts).encode())


# =========================================================================== driver


def _select(only: str | None) -> list[Fixture]:
    if not only:
        return list(REGISTRY)
    return [f for f in REGISTRY if fnmatch.fnmatch(f.id, only)]


def generate(out_root: Path, fixtures: Iterable[Fixture], *, edge: bool) -> list[Path]:
    written: list[Path] = []
    for fx in fixtures:
        if fx.on_demand and not edge:
            continue
        target = out_root / fx.relpath
        target.parent.mkdir(parents=True, exist_ok=True)
        fx.build(target, fx)
        written.append(target)
    return written


def _digest(path: Path) -> str:
    if path.is_symlink():
        return "symlink:" + os.readlink(path)
    h = hashlib.sha256()
    size = path.stat().st_size
    with open(path, "rb") as f:
        if size > 8 * 1024 * 1024:  # sparse/on-demand: size + first MiB is enough
            h.update(f.read(1024 * 1024))
            return f"size={size}:head={h.hexdigest()}"
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def check(fixtures: list[Fixture], *, edge: bool) -> int:
    with tempfile.TemporaryDirectory(prefix="evoq-corpus-check-") as tmp:
        tmp_root = Path(tmp)
        generate(tmp_root, fixtures, edge=edge)
        failures = 0
        for fx in fixtures:
            if fx.on_demand and not edge:
                continue
            committed = CORPUS_ROOT / fx.relpath
            fresh = tmp_root / fx.relpath
            if not committed.exists() and not committed.is_symlink():
                print(f"MISSING   {fx.id}: {fx.relpath} (run without --check first)")
                failures += 1
                continue
            if _digest(committed) != _digest(fresh):
                print(f"DRIFT     {fx.id}: {fx.relpath}")
                failures += 1
            else:
                print(f"identical {fx.id}: {fx.relpath}")
        return failures


def size_report(fixtures: list[Fixture]) -> int:
    over = 0
    for fx in fixtures:
        p = CORPUS_ROOT / fx.relpath
        if fx.on_demand or fx.symlink_to or not p.exists():
            continue
        size = p.stat().st_size
        if size > MAX_COMMITTED_BYTES:
            print(f"TOO LARGE {fx.id}: {size:,} bytes > {MAX_COMMITTED_BYTES:,}")
            over += 1
    return over


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument(
        "--out",
        type=Path,
        default=CORPUS_ROOT,
        help="corpus root to write into (default: this directory)",
    )
    ap.add_argument(
        "--check",
        action="store_true",
        help="regenerate to a temp dir and diff against the committed files",
    )
    ap.add_argument(
        "--edge",
        action="store_true",
        help="also create on-demand fixtures (EDGE-01 sparse 201 MiB)",
    )
    ap.add_argument("--only", help="glob over fixture ids, e.g. 'PDF-*'")
    ap.add_argument("--list", action="store_true", help="print the registry and exit")
    args = ap.parse_args(argv)

    fixtures = _select(args.only)
    if not fixtures:
        print(f"no fixtures match {args.only!r}", file=sys.stderr)
        return 2
    if args.list:
        for fx in fixtures:
            flag = (
                " (on-demand)"
                if fx.on_demand
                else (" (symlink)" if fx.symlink_to else "")
            )
            print(f"{fx.id:8} {fx.relpath}{flag}")
        return 0
    if args.check:
        failures = check(fixtures, edge=args.edge)
        failures += size_report(fixtures)
        print(
            "--check:",
            "OK (byte-identical)" if failures == 0 else f"{failures} problem(s)",
        )
        return 1 if failures else 0

    written = generate(args.out, fixtures, edge=args.edge)
    for p in written:
        if p.is_symlink():
            print(f"wrote {p.relative_to(args.out)} -> {os.readlink(p)}")
        else:
            print(f"wrote {p.relative_to(args.out)} ({p.stat().st_size:,} bytes)")
    over = size_report(fixtures) if args.out == CORPUS_ROOT else 0
    return 1 if over else 0


if __name__ == "__main__":
    sys.exit(main())
