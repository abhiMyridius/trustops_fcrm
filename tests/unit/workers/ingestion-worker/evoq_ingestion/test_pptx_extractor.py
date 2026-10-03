"""Tests for ``extractors/pptx_extractor.py`` — the model-based ``PptxExtractor``.

Corpus fixtures (lane L9, ``AEO/contracts/fixtures/corpus/generated/PPTX-*``)
are cross-checked against ``scenarios.yaml``; hand-built ``python-pptx``
fixtures cover cases the corpus doesn't isolate (a nested group-in-group,
notes gating, a merged-cell table, a chart, a picture).
"""

from __future__ import annotations

import dataclasses
import sys
import types
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _bootstrap_evoq_ingestion_package() -> None:
    """Make the package importable despite an unrelated, in-flight lane's breakage.

    ``evoq_ingestion/__init__.py`` eagerly imports the whole tree, which as
    of this run fails deep inside modules **outside this lane's scope**
    (``extractors/{pptx,text}_extractor.py`` only): ``sources/jira_source.py``
    still imports ``flatten_adf_to_text`` from ``core/utils.py`` after a
    concurrent lane deleted it there, and ``extractors/registry.py`` still
    imports the pre-model ``PdfExtractor`` name after a concurrent lane moved
    ``pdf_extractor.py`` to the new ``Document`` model. Neither module is
    reachable from the PPTX/TXT extractors under test. If the real package
    already imports cleanly (i.e. those lanes have since landed), this is a
    no-op; otherwise a lightweight stand-in for the top package — pointed at
    the real package directory via ``__path__`` — lets every submodule this
    test needs still import normally straight off disk, without ever
    executing the eager (and currently broken) top-level ``__init__.py``.
    """
    if "evoq_ingestion" in sys.modules:
        return
    try:
        import evoq_ingestion  # noqa: F401
    except ImportError:
        stub = types.ModuleType("evoq_ingestion")
        stub.__path__ = [str(Path(__file__).resolve().parents[1])]
        sys.modules["evoq_ingestion"] = stub


_bootstrap_evoq_ingestion_package()

from evoq_ingestion.core.codes import FailureCode  # noqa: E402
from evoq_ingestion.core.config import ExtractionOptions  # noqa: E402
from evoq_ingestion.core.limits import Deadline, Limits  # noqa: E402
from evoq_ingestion.extractors.base import CorruptFile, LimitExceeded  # noqa: E402
from evoq_ingestion.extractors.gate import FileFacts, FileGate, FileKind  # noqa: E402
from evoq_ingestion.extractors.pptx_extractor import PptxExtractor  # noqa: E402
from evoq_ingestion.model.blocks import (  # noqa: E402
    Heading,
    ImageRef,
    ListBlock,
    Note,
    NoteKind,
    Paragraph,
)
from evoq_ingestion.model.document import Document  # noqa: E402
from evoq_ingestion.model.renderer import GfmRenderer  # noqa: E402

CORPUS_DIR = Path(__file__).resolve().parent / "fixtures" / "contracts" / "corpus"
SCENARIOS_PATH = CORPUS_DIR / "scenarios.yaml"


def _options(**overrides: object) -> ExtractionOptions:
    opts = ExtractionOptions()
    if overrides:
        opts = dataclasses.replace(opts, **overrides)  # type: ignore[arg-type]
    return opts


async def _extract(path: Path, options: ExtractionOptions | None = None) -> Document:
    """Full pipeline slice: gate inspection then the extractor, as production runs it."""
    facts = FileGate().inspect(path)
    extractor = PptxExtractor(options or _options())
    return await extractor.extract(path, facts, Deadline.unlimited())


def _fake_facts(path: Path) -> FileFacts:
    """A ``FileFacts`` bypassing the gate, for extractor-only edge-case tests."""
    return FileFacts(
        path=path, size=path.stat().st_size, sha256="x", ext=".pptx", kind=FileKind.pptx
    )


def _all_text(doc: Document) -> str:
    return GfmRenderer().render(doc).markdown


def _load_scenario(scenario_id: str) -> dict:
    data = yaml.safe_load(SCENARIOS_PATH.read_text())
    for item in data["scenarios"]:
        if item["id"] == scenario_id:
            return item
    raise AssertionError(f"scenario {scenario_id} not found")


# --------------------------------------------------------------------- corpus
@pytest.mark.skipif(not CORPUS_DIR.exists(), reason="corpus fixtures not present")
class TestCorpusFixtures:
    async def test_pptx01_six_slides_notes(self):
        scenario = _load_scenario("PPTX-01")
        path = CORPUS_DIR / scenario["file"]
        doc = await _extract(path)

        assert not doc.units.failed
        assert doc.units.slide_count == 6
        assert len(doc.slide_breaks()) == 6
        assert all(not b.hidden for b in doc.slide_breaks())
        assert not doc.tables()

        body = _all_text(doc)
        for expected in scenario["expected"]["canonical"]["body"]["contains"]:
            assert expected in body, expected
        assert "--- Slide 1 ---" in body
        assert "--- Slide 6 ---" in body

    async def test_pptx02_table_group_hidden(self):
        scenario = _load_scenario("PPTX-02")
        path = CORPUS_DIR / scenario["file"]
        doc = await _extract(path)

        assert doc.units.slide_count == 4
        assert len(doc.tables()) == 1
        slide_breaks = doc.slide_breaks()
        assert len(slide_breaks) == 4
        hidden = [b for b in slide_breaks if b.hidden]
        assert len(hidden) == 1
        assert hidden[0].n == 4
        assert any(w.startswith("hidden_slide") for w in doc.warnings)

        body = _all_text(doc)
        for expected in scenario["expected"]["canonical"]["body"]["contains"]:
            assert expected in body, expected

        table = doc.tables()[0]
        header_texts = [c.text for c in table.rows[0].cells]
        assert header_texts == ["Service", "Owner", "Tier"]

    async def test_pptx03_slide_limit_101(self):
        scenario = _load_scenario("PPTX-03")
        path = CORPUS_DIR / scenario["file"]
        with pytest.raises(LimitExceeded) as exc_info:
            await _extract(path)
        assert exc_info.value.failure_code is FailureCode.slide_limit_exceeded
        assert scenario["expected"]["canonical"]["status"] == "FAILED"


# ------------------------------------------------------------------- fixtures
def _new_prs() -> Presentation:
    return Presentation()


def _add_title_slide(prs: Presentation, title: str, body: str | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    if body is not None:
        slide.placeholders[1].text_frame.text = body


class TestGroupedShapes:
    async def test_nested_group_in_group_survives(self, tmp_path: Path):
        prs = _new_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box1 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        box1.text_frame.text = "outer box text"
        box2 = slide.shapes.add_textbox(Inches(3), Inches(1), Inches(2), Inches(1))
        box2.text_frame.text = "DEEP-NESTED-TEXT"
        inner_group = slide.shapes.add_group_shape([box2])
        outer_group = slide.shapes.add_group_shape([box1, inner_group])
        assert outer_group is not None
        path = tmp_path / "nested_group.pptx"
        prs.save(path)

        doc = await _extract(path)
        body = _all_text(doc)
        assert "outer box text" in body
        assert "DEEP-NESTED-TEXT" in body


class TestTables:
    async def test_merged_cells_render_span_not_duplicate(self, tmp_path: Path):
        prs = _new_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        gframe = slide.shapes.add_table(
            3, 3, Inches(1), Inches(1), Inches(5), Inches(3)
        )
        table = gframe.table
        # Merge first, then type text — matches how PowerPoint itself produces
        # a merged cell (typing into two cells and merging them concatenates
        # both paragraphs into the origin, which is python-pptx's own
        # ``merge()`` behaviour, not something this extractor should mask).
        table.cell(0, 0).merge(table.cell(1, 0))  # vertical: column 0, rows 0-1
        table.cell(2, 1).merge(table.cell(2, 2))  # horizontal: row 2, columns 1-2
        for r in range(3):
            for c in range(3):
                if table.cell(r, c).is_spanned:
                    continue
                table.cell(r, c).text = f"r{r}c{c}"
        path = tmp_path / "merged_table.pptx"
        prs.save(path)

        doc = await _extract(path)
        tables = doc.tables()
        assert len(tables) == 1
        table_block = tables[0]

        row0_cells = table_block.rows[0].cells
        origin = row0_cells[0]
        assert origin.text == "r0c0"
        assert origin.rowspan == 2
        # Row 0 should have 3 cells total (origin + 2 plain), none duplicated.
        assert len(row0_cells) == 3

        row1_cells = table_block.rows[1].cells
        assert [c.text for c in row1_cells] == ["r1c1", "r1c2"]

        row2_cells = table_block.rows[2].cells
        assert row2_cells[0].text == "r2c0"
        merged = row2_cells[1]
        assert merged.text == "r2c1"
        assert merged.colspan == 2
        assert len(row2_cells) == 2

        # The renderer's default rowspan policy *repeats* the origin text on
        # the covered row (a rendering policy, not extractor duplication);
        # colspan instead pads with an empty cell — never repeats text.
        rendered = GfmRenderer().render(doc).markdown
        assert rendered.count("r0c0") == 2
        assert rendered.count("r2c1") == 1


class TestNotes:
    async def test_notes_included_when_enabled(self, tmp_path: Path):
        prs = _new_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(2), Inches(1)
        ).text_frame.text = "content"
        slide.notes_slide.notes_text_frame.text = "SPEAKER-NOTE-TEXT"
        path = tmp_path / "notes.pptx"
        prs.save(path)

        doc = await _extract(path, _options(pptx_include_notes=True))
        notes = [
            b
            for b in doc.iter_blocks()
            if isinstance(b, Note) and b.kind is NoteKind.speaker_notes
        ]
        assert len(notes) == 1
        assert "SPEAKER-NOTE-TEXT" in notes[0].text

    async def test_notes_excluded_when_disabled(self, tmp_path: Path):
        prs = _new_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(2), Inches(1)
        ).text_frame.text = "content"
        slide.notes_slide.notes_text_frame.text = "SHOULD-NOT-APPEAR"
        path = tmp_path / "notes_off.pptx"
        prs.save(path)

        doc = await _extract(path, _options(pptx_include_notes=False))
        assert "SHOULD-NOT-APPEAR" not in _all_text(doc)
        assert not any(
            isinstance(b, Note) and b.kind is NoteKind.speaker_notes
            for b in doc.iter_blocks()
        )


class TestHiddenSlide:
    async def test_hidden_slide_flagged_but_content_extracted(self, tmp_path: Path):
        prs = _new_prs()
        visible = prs.slides.add_slide(prs.slide_layouts[6])
        visible.shapes.add_textbox(
            Inches(1), Inches(1), Inches(2), Inches(1)
        ).text_frame.text = "visible"
        hidden_slide = prs.slides.add_slide(prs.slide_layouts[6])
        hidden_slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(2), Inches(1)
        ).text_frame.text = "HIDDEN-CONTENT"
        hidden_slide._element.set("show", "0")  # noqa: SLF001 - test setup, mirrors PowerPoint's own attr
        path = tmp_path / "hidden.pptx"
        prs.save(path)

        doc = await _extract(path)
        breaks = doc.slide_breaks()
        assert breaks[0].hidden is False
        assert breaks[1].hidden is True
        assert any(w == "hidden_slide:2" for w in doc.warnings)
        assert "HIDDEN-CONTENT" in _all_text(doc)


class TestSlideLimit:
    async def test_101_slides_raises_before_parsing(self, tmp_path: Path):
        prs = _new_prs()
        for i in range(101):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(2), Inches(1)
            ).text_frame.text = f"s{i}"
        path = tmp_path / "many_slides.pptx"
        prs.save(path)

        with pytest.raises(LimitExceeded) as exc_info:
            await _extract(path, _options(limits=Limits(max_slides=100)))
        assert exc_info.value.failure_code is FailureCode.slide_limit_exceeded


class TestChart:
    async def test_chart_becomes_summary_note_with_warning(self, tmp_path: Path):
        prs = _new_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        chart_data = CategoryChartData()
        chart_data.categories = ["a", "b"]
        chart_data.add_series("Revenue", (1, 2))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(4),
            Inches(3),
            chart_data,
        )
        path = tmp_path / "chart.pptx"
        prs.save(path)

        doc = await _extract(path)
        chart_notes = [
            b
            for b in doc.iter_blocks()
            if isinstance(b, Note) and b.kind is NoteKind.chart
        ]
        assert len(chart_notes) == 1
        assert "Revenue" in chart_notes[0].text
        assert any(w.startswith("chart_summarised") for w in doc.warnings)


class TestPicture:
    async def test_picture_becomes_image_ref(self, tmp_path: Path):
        from PIL import Image

        img_path = tmp_path / "pixel.png"
        Image.new("RGB", (2, 2), color="red").save(img_path)

        prs = _new_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(
            str(img_path), Inches(1), Inches(1), Inches(1), Inches(1)
        )
        path = tmp_path / "picture.pptx"
        prs.save(path)

        doc = await _extract(path)
        images = doc.image_refs()
        assert len(images) == 1
        assert isinstance(images[0], ImageRef)
        assert images[0].target.startswith("attachment:")
        assert doc.units.image_count == 1


class TestTitleAndParagraphs:
    async def test_title_becomes_heading_and_single_paragraph_stays_plain(
        self, tmp_path: Path
    ):
        prs = _new_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "My Slide Title"
        slide.placeholders[1].text_frame.text = "Just one line of body text."
        path = tmp_path / "single_para.pptx"
        prs.save(path)

        doc = await _extract(path)
        headings = doc.headings()
        assert len(headings) == 1
        assert headings[0] == Heading(2, text="My Slide Title")

        paragraphs = [b for b in doc.iter_blocks() if isinstance(b, Paragraph)]
        assert any(p.text == "Just one line of body text." for p in paragraphs)

    async def test_multi_paragraph_body_becomes_list(self, tmp_path: Path):
        prs = _new_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Bulleted"
        tf = slide.placeholders[1].text_frame
        tf.text = "first bullet"
        p2 = tf.add_paragraph()
        p2.text = "second bullet"
        path = tmp_path / "bulleted.pptx"
        prs.save(path)

        doc = await _extract(path)
        lists = [b for b in doc.iter_blocks() if isinstance(b, ListBlock)]
        assert len(lists) == 1
        item_texts = [item.text for item in lists[0].items]
        assert item_texts == ["first bullet", "second bullet"]


class TestCorruptFile:
    async def test_non_zip_pptx_raises_corrupt_file(self, tmp_path: Path):
        # Bypass the gate here: plain-text bytes named ``.pptx`` are correctly
        # refused by ``FileGate`` as ``extension_mismatch`` (its own, more
        # specific diagnosis) before the extractor is ever reached. This test
        # targets the extractor's *own* ``classify_error`` mapping — a
        # zip-shaped-but-unparseable-by-python-pptx container — so it drives
        # the extractor directly against a fake, gate-bypassing ``FileFacts``.
        path = tmp_path / "bad.pptx"
        path.write_bytes(b"PK\x03\x04 not actually a valid OOXML zip stream")
        extractor = PptxExtractor(_options())
        with pytest.raises(CorruptFile):
            await extractor.extract(path, _fake_facts(path), Deadline.unlimited())


class TestPerUnitDegradation:
    async def test_one_bad_slide_does_not_kill_the_others(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        prs = _new_prs()
        for i in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(2), Inches(1)
            ).text_frame.text = f"ok-{i}"
        path = tmp_path / "three_slides.pptx"
        prs.save(path)

        extractor = PptxExtractor(_options())
        original = extractor._parse_slide

        def boom(doc, slide, idx, *, include_notes):
            if idx == 2:
                raise RuntimeError("simulated slide failure")
            return original(doc, slide, idx, include_notes=include_notes)

        monkeypatch.setattr(extractor, "_parse_slide", boom)
        facts = _fake_facts(path)
        doc = await extractor.extract(path, facts, Deadline.unlimited())

        # 1-based slide idx=2 is the *second* slide, holding the i=1 textbox
        # ("ok-1"); it is the one ``boom`` fails.
        assert len(doc.units.failed) == 1
        assert doc.units.failed[0].ref == "2"
        assert doc.units.slide_count == 3
        assert len(doc.slide_breaks()) == 3  # markers still emitted for every slide
        body = _all_text(doc)
        assert "ok-0" in body
        assert "ok-1" not in body
        assert "ok-2" in body
